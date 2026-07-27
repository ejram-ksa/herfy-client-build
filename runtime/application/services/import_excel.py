from __future__ import annotations
import csv
import logging
import re
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from runtime.shared.errors import SERVICE_OPERATION_EXCEPTIONS
from runtime.shared.materials import normalize_material_code
from runtime.shared.numbers import parse_excel_number

logger = logging.getLogger(__name__)
MAX_IMPORT_FILE_BYTES = 50 * 1024 * 1024
MAX_IMPORT_ROWS = 100_000
MAX_IMPORT_COLUMNS = 256
MAX_IMPORT_CELLS = 2_000_000
_SKIP_CODES = {"total", "grand total", "subtotal", "المجموع", "اجمالي", "الإجمالي"}
_LABEL_ALIASES = {
    "material": {
        "material",
        "material code",
        "material number",
        "material no",
        "رمز الصنف",
        "رقم الصنف",
        "رقم المادة",
    },
    "item": {"item", "code", "item code", "item no", "المادة", "الصنف"},
    "name": {
        "description",
        "material description",
        "item description",
        "name",
        "material name",
        "اسم الصنف",
        "اسم المادة",
        "الوصف",
        "وصف الصنف",
    },
    "receipts": {
        "receipts",
        "rec qty",
        "rec. qty",
        "qty received",
        "received",
        "receipt qty",
        "receipt quantity",
        "quantity received",
        "receipt",
        "الاستلام",
        "المستلم",
        "كمية الاستلام",
        "الكمية المستلمة",
    },
    "begining": {
        "begining",
        "beginning",
        "begin",
        "opening",
        "opening qty",
        "qty counted",
        "counted qty",
        "counted",
        "opening balance",
        "الرصيد الافتتاحي",
        "الافتتاحي",
        "الجرد",
        "كمية الجرد",
        "الكمية المعدودة",
    },
    "uom": {"uom", "unit", "unit of measure", "u/m", "الوحدة"},
    "qty": {"qty", "quantity", "الكمية"},
}


@dataclass
class SimpleTable:
    columns: list[Any]
    rows: list[dict[Any, Any]]

    @property
    def shape(self) -> tuple[int, int]:
        return (len(self.rows), len(self.columns))

    def iterrows(self):
        yield from enumerate(self.rows)


def _normalize_label(value) -> str:
    s = str(value or "").strip().lower()
    s = re.sub("\\s+", " ", s)
    s = s.replace("_", " ").replace("-", " ")
    s = s.replace(".", "").replace(":", "")
    for canonical, aliases in _LABEL_ALIASES.items():
        if s == canonical or s in aliases:
            return canonical
    return s


def _cell(row, col):
    try:
        if row is None or col is None:
            return None
        if isinstance(row, dict):
            return row.get(col)
        if isinstance(row, list | tuple) and isinstance(col, int):
            return row[col] if 0 <= col < len(row) else None
        return row[col]
    except SERVICE_OPERATION_EXCEPTIONS:
        return None


def _validate_import_file(path: Path) -> Path:
    resolved = path.expanduser().resolve()
    if not resolved.is_file():
        raise ValueError(f"Import file does not exist: {resolved}")
    size = resolved.stat().st_size
    if size > MAX_IMPORT_FILE_BYTES:
        raise ValueError(
            f"Import file is too large: {size} bytes; "
            f"maximum={MAX_IMPORT_FILE_BYTES}"
        )
    return resolved


def _bounded_row(values: Iterable[Any], row_number: int) -> list[Any]:
    row = list(values)
    if len(row) > MAX_IMPORT_COLUMNS:
        raise ValueError(
            f"Import row {row_number} has too many columns: "
            f"{len(row)}; maximum={MAX_IMPORT_COLUMNS}"
        )
    return row


def _append_bounded_row(
    rows: list[list[Any]],
    values: Iterable[Any],
    row_number: int,
    cell_count: int,
) -> int:
    if row_number > MAX_IMPORT_ROWS:
        raise ValueError(
            f"Import contains too many rows: maximum={MAX_IMPORT_ROWS}"
        )
    row = _bounded_row(values, row_number)
    next_cell_count = cell_count + len(row)
    if next_cell_count > MAX_IMPORT_CELLS:
        raise ValueError(
            f"Import contains too many cells: maximum={MAX_IMPORT_CELLS}"
        )
    rows.append(row)
    return next_cell_count


def _read_csv_rows(path: Path) -> list[list[Any]]:
    candidates = ("utf-8-sig", "utf-8", "cp1256", "cp1252")
    last_error: Exception | None = None
    for enc in candidates:
        try:
            with open(path, encoding=enc, newline="") as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample or ",")
                except SERVICE_OPERATION_EXCEPTIONS:
                    dialect = csv.excel
                rows: list[list[Any]] = []
                cell_count = 0
                for row_number, row in enumerate(csv.reader(fh, dialect), 1):
                    cell_count = _append_bounded_row(
                        rows, row, row_number, cell_count
                    )
                return rows
        except UnicodeDecodeError as exc:
            last_error = exc
    if last_error is not None:
        raise last_error
    return []


def _read_xlsx_rows(path: Path) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        ws = wb.active
        if int(ws.max_row or 0) > MAX_IMPORT_ROWS:
            raise ValueError(
                f"Import contains too many rows: {ws.max_row}; "
                f"maximum={MAX_IMPORT_ROWS}"
            )
        if int(ws.max_column or 0) > MAX_IMPORT_COLUMNS:
            raise ValueError(
                f"Import contains too many columns: {ws.max_column}; "
                f"maximum={MAX_IMPORT_COLUMNS}"
            )
        if int(ws.max_row or 0) * int(ws.max_column or 0) > MAX_IMPORT_CELLS:
            raise ValueError(
                f"Import contains too many cells: maximum={MAX_IMPORT_CELLS}"
            )
        rows: list[list[Any]] = []
        cell_count = 0
        for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
            cell_count = _append_bounded_row(
                rows, row, row_number, cell_count
            )
        return rows
    finally:
        wb.close()


def _read_xls_rows(path: Path) -> list[list[Any]]:
    import xlrd

    book = xlrd.open_workbook(str(path), on_demand=True)
    try:
        sheet = book.sheet_by_index(0)
        if int(sheet.nrows) > MAX_IMPORT_ROWS:
            raise ValueError(
                f"Import contains too many rows: {sheet.nrows}; "
                f"maximum={MAX_IMPORT_ROWS}"
            )
        if int(sheet.ncols) > MAX_IMPORT_COLUMNS:
            raise ValueError(
                f"Import contains too many columns: {sheet.ncols}; "
                f"maximum={MAX_IMPORT_COLUMNS}"
            )
        if int(sheet.nrows) * int(sheet.ncols) > MAX_IMPORT_CELLS:
            raise ValueError(
                f"Import contains too many cells: maximum={MAX_IMPORT_CELLS}"
            )
        rows: list[list[Any]] = []
        cell_count = 0
        for row_number in range(1, sheet.nrows + 1):
            cell_count = _append_bounded_row(
                rows,
                sheet.row_values(row_number - 1),
                row_number,
                cell_count,
            )
        return rows
    finally:
        release = getattr(book, "release_resources", None)
        if callable(release):
            release()


def _read_table_rows(path: str | Path) -> list[list[Any]]:
    file_path = _validate_import_file(Path(path))
    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _read_xlsx_rows(file_path)
    if suffix == ".xls":
        return _read_xls_rows(file_path)
    if suffix in {".csv", ".txt", ".tsv"}:
        return _read_csv_rows(file_path)
    raise ValueError(
        f"Unsupported import file type: {suffix or '<no extension>'}"
    )


def _unique_columns(values: Iterable[Any]) -> list[Any]:
    result: list[Any] = []
    seen: dict[str, int] = {}
    for idx, raw in enumerate(values):
        base = str(raw).strip() if raw not in (None, "") else f"column_{idx + 1}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        result.append(base if count == 1 else f"{base}_{count}")
    return result


def read_smart_excel_en(path: str):
    raw_rows = _read_table_rows(path)
    if not raw_rows:
        return SimpleTable(columns=[], rows=[])
    header_row = -1
    keys = [
        "material",
        "item",
        "code",
        "uom",
        "unit",
        "description",
        "name",
        "qty",
        "receipts",
        "qty counted",
        "counted qty",
        "begin",
        "begining",
    ]
    limit = min(80, len(raw_rows))
    for i in range(limit):
        row_text = " | ".join([_normalize_label(x) for x in raw_rows[i]])
        if any((k in row_text for k in keys)):
            header_row = i
            break
    if header_row >= 0:
        columns = _unique_columns(raw_rows[header_row])
        body_rows = raw_rows[header_row + 1 :]
    else:
        width = max((len(r) for r in raw_rows), default=0)
        columns = list(range(width))
        body_rows = raw_rows
    normalized_rows: list[dict[Any, Any]] = []
    for raw in body_rows:
        row_values = list(raw)
        if len(row_values) < len(columns):
            row_values.extend([None] * (len(columns) - len(row_values)))
        elif len(row_values) > len(columns):
            row_values = row_values[: len(columns)]
        normalized_rows.append(
            {col: row_values[idx] for idx, col in enumerate(columns)}
        )
    return SimpleTable(columns=list(columns), rows=normalized_rows)


def _norm(df):
    return {c: _normalize_label(c) for c in df.columns}


def _first_matching_label(normalized_cols, labels):
    for want in labels:
        for column, label in normalized_cols.items():
            if label == want:
                return column
    return None


def _pick(colsmap, *names):
    return _first_matching_label(colsmap, names)


def _pick_preferred(colsmap, preferred_labels, fallback_labels=()):
    normalized = {c: str(lbl).strip().lower() for c, lbl in colsmap.items()}
    return _first_matching_label(normalized, preferred_labels) or _first_matching_label(
        normalized, fallback_labels
    )


def _pick_column_or_position(
    colsmap, df, *, preferred_labels, fallback_labels=(), default_index: int = 0
):
    col = _pick_preferred(
        colsmap, preferred_labels=preferred_labels, fallback_labels=fallback_labels
    )
    if col is not None:
        return col
    return df.columns[default_index] if df.shape[1] > default_index else None


def _pick_material_col(colsmap, df):
    return _pick_column_or_position(
        colsmap,
        df,
        preferred_labels=(
            "material",
            "material number",
            "material no",
            "material code",
            "رقم المادة",
            "رقم الصنف",
            "رمز الصنف",
        ),
        fallback_labels=("item code", "item no", "code", "item", "المادة", "الصنف"),
        default_index=0,
    )


def _iter_valid_material_rows(df, material_col):
    for _, row in df.iterrows():
        code = normalize_material_code(_cell(row, material_col))
        if not code or code.strip().lower() in _SKIP_CODES:
            continue
        yield (row, code)


def _pick_name_col(colsmap, df):
    return _pick_column_or_position(
        colsmap,
        df,
        preferred_labels=(
            "material description",
            "description",
            "material name",
            "item description",
            "name",
            "اسم المادة",
            "اسم الصنف",
            "الوصف",
        ),
        default_index=1,
    )


def _pick_begin_col(colsmap, df):
    return _pick_column_or_position(
        colsmap,
        df,
        preferred_labels=(
            "begining",
            "beginning",
            "qty counted",
            "counted qty",
            "counted",
            "opening qty",
            "opening balance",
            "الجرد",
            "كمية الجرد",
            "الكمية المعدودة",
        ),
        fallback_labels=("qty", "quantity", "الكمية"),
        default_index=2,
    )


def _pick_receipts_col(colsmap, df):
    return _pick_column_or_position(
        colsmap,
        df,
        preferred_labels=(
            "receipts",
            "rec qty",
            "rec. qty",
            "qty received",
            "quantity received",
            "receipt qty",
            "received",
            "receipt",
            "الاستلام",
            "كمية الاستلام",
            "الكمية المستلمة",
        ),
        fallback_labels=("qty", "quantity", "الكمية"),
        default_index=2,
    )


def parse_receipts_uom_from_colD(path: str):
    df = read_smart_excel_en(path)
    cols = _norm(df)
    mat_col = _pick_material_col(cols, df)
    name_col = _pick_name_col(cols, df)
    rec_col = _pick_receipts_col(cols, df)
    uom_col = _pick(cols, "uom") or (df.columns[3] if df.shape[1] >= 4 else None)
    out = {}
    uom_map = {}
    for row, code in _iter_valid_material_rows(df, mat_col):
        name = str(_cell(row, name_col) or "").strip()
        rec = parse_excel_number(_cell(row, rec_col))
        uom = str(_cell(row, uom_col) or "").strip()
        bucket = out.setdefault(
            code, {"material": code, "name": name, "uom": uom, "receipts": 0.0}
        )
        bucket.update(receipts=float(bucket.get("receipts", 0.0)) + rec)
        if uom and (not bucket.get("uom")):
            bucket.update(uom=uom)
        if name and (not bucket.get("name")):
            bucket.update(name=name)
        if uom:
            uom_map[code] = uom
    return (out, uom_map)


def parse_begining(path: str):
    df = read_smart_excel_en(path)
    cols = _norm(df)
    c_mat = _pick_material_col(cols, df)
    c_qty = _pick_begin_col(cols, df)
    c_name = _pick_name_col(cols, df)
    tmp = {}
    for row, code in _iter_valid_material_rows(df, c_mat):
        q = parse_excel_number(_cell(row, c_qty))
        name = str(_cell(row, c_name) or "").strip()
        tmp.setdefault(code, []).append({"qty": q, "name": name})
    return {k: arr[-1] for k, arr in tmp.items()}
